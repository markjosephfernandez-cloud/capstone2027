import os
import re
import sys
from pathlib import Path

DEPS = Path(__file__).resolve().parent / "docx-deps"
sys.path.insert(0, str(DEPS))

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "writable" / "Revised_Chapters_1_and_2.md"
OUTPUT = ROOT / "Revised_Chapter_1_and_2_Greyhound_Hub_2.0.docx"


def set_font(run, name="Times New Roman", size=12, bold=None, italic=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Inches(0.5)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in (
        ("Title", 16, 0, 12),
        ("Subtitle", 14, 0, 12),
        ("Heading 1", 14, 12, 12),
        ("Heading 2", 12, 12, 6),
        ("Heading 3", 12, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Times New Roman"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.first_line_indent = Inches(0)
        style.paragraph_format.line_spacing = 1.0
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Academic List" not in [s.name for s in doc.styles]:
        academic_list = doc.styles.add_style("Academic List", WD_STYLE_TYPE.PARAGRAPH)
    else:
        academic_list = doc.styles["Academic List"]
    academic_list.base_style = doc.styles["Normal"]
    academic_list.paragraph_format.left_indent = Inches(0.5)
    academic_list.paragraph_format.first_line_indent = Inches(-0.25)
    academic_list.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    academic_list.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    academic_list.paragraph_format.space_after = Pt(4)

    if "Reference" not in [s.name for s in doc.styles]:
        reference = doc.styles.add_style("Reference", WD_STYLE_TYPE.PARAGRAPH)
    else:
        reference = doc.styles["Reference"]
    reference.base_style = doc.styles["Normal"]
    reference.paragraph_format.left_indent = Inches(0.5)
    reference.paragraph_format.first_line_indent = Inches(-0.5)
    reference.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    reference.paragraph_format.line_spacing = 1.0
    reference.paragraph_format.space_after = Pt(8)


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    set_font(run, size=10)
    fld_char_1 = OxmlElement("w:fldChar")
    fld_char_1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char_2 = OxmlElement("w:fldChar")
    fld_char_2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char_1, instr_text, fld_char_2])


def add_rich_text(paragraph, text, base_size=12):
    # Supports the small subset of Markdown used in the source: **bold** and *italics*.
    pattern = re.compile(r"(\*\*.+?\*\*|\*[^*]+?\*)")
    pos = 0
    for match in pattern.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_font(run, size=base_size)
        token = match.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_font(run, size=base_size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_font(run, size=base_size, italic=True)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_font(run, size=base_size)


def add_heading(doc, text, level):
    if level == 1:
        if text == "GREYHOUND HUB 2.0":
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(72)
            add_rich_text(p, text, 16)
            for run in p.runs:
                run.bold = True
            return
        p = doc.add_paragraph(style="Heading 1")
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.page_break_before = True
        p.paragraph_format.space_before = Pt(0)
        add_rich_text(p, text, 14)
        for run in p.runs:
            run.bold = True
        return

    style = "Heading 2" if level == 2 else "Heading 3"
    p = doc.add_paragraph(style=style)
    if text in {"REVISED CHAPTERS 1 AND 2", "THE PROBLEM AND ITS BACKGROUND", "REVIEW OF RELATED LITERATURE AND STUDIES"}:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_rich_text(p, text, 14 if level == 2 and text == "REVISED CHAPTERS 1 AND 2" else 12)
    for run in p.runs:
        run.bold = True


def add_body_paragraph(doc, text, in_references=False):
    if re.match(r"^\d+\.\s", text):
        p = doc.add_paragraph(style="Academic List")
        add_rich_text(p, text)
        return

    if in_references:
        p = doc.add_paragraph(style="Reference")
        add_rich_text(p, text)
        return

    p = doc.add_paragraph(style="Normal")
    if text in {"Foundation University", "Working research manuscript based on the current implemented system"}:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.first_line_indent = Inches(0)
        p.paragraph_format.line_spacing = 1.0
        p.paragraph_format.space_after = Pt(10)
        add_rich_text(p, text)
        return

    if text.startswith("**") and ".**" in text:
        p.paragraph_format.first_line_indent = Inches(0)
    add_rich_text(p, text)


def build_document():
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1.5)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.4)
    section.footer_distance = Inches(0.5)
    configure_styles(doc)
    add_page_number(section.footer.paragraphs[0])

    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    in_references = False
    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue
        if line == "---PAGEBREAK---":
            doc.add_page_break()
            continue
        if line.startswith("### "):
            add_heading(doc, line[4:], 3)
            continue
        if line.startswith("## "):
            heading = line[3:]
            add_heading(doc, heading, 2)
            in_references = heading == "REFERENCES"
            continue
        if line.startswith("# "):
            add_heading(doc, line[2:], 1)
            in_references = False
            continue
        add_body_paragraph(doc, line, in_references=in_references)

    props = doc.core_properties
    props.title = "Greyhound Hub 2.0 — Revised Chapters 1 and 2"
    props.subject = "Revised capstone research manuscript aligned with the current system"
    props.author = "Foundation University Research Team"
    props.keywords = "Greyhound Hub 2.0, digital asset management, AI, media passport, Foundation University"
    props.comments = "Generated from the audited current system implementation."

    doc.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    output = build_document()
    print(output)
